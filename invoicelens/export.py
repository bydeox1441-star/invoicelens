"""Write results as Excel (formatted, with a review column), CSV or JSON."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import HEADER_FIELDS, Result

REVIEW_FILL = PatternFill("solid", start_color="FFF4CE")
HEADER_FILL = PatternFill("solid", start_color="1B2A41")


def invoices_frame(results: list[Result]) -> pd.DataFrame:
    rows = []
    for r in results:
        row = {"file": r.file}
        row.update({f: getattr(r.invoice, f) for f in HEADER_FIELDS})
        row.update({
            "line_count": len(r.invoice.line_items),
            "method": r.method,
            "layout": r.layout,
            "scanned": r.scanned,
            "ocr_confidence": r.ocr_confidence,
            "needs_review": r.needs_review,
            "review_reason": r.review_reason,
        })
        rows.append(row)
    return pd.DataFrame(rows)


def lines_frame(results: list[Result]) -> pd.DataFrame:
    rows = [
        {"file": r.file, "invoice_number": r.invoice.invoice_number, "line": n,
         "description": li.description, "quantity": li.quantity, "unit_price": li.unit_price, "amount": li.amount}
        for r in results for n, li in enumerate(r.invoice.line_items, start=1)
    ]
    return pd.DataFrame(rows, columns=["file", "invoice_number", "line", "description", "quantity", "unit_price", "amount"])


def _style_sheet(ws, review_col: int | None = None) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)
    if review_col:
        for row in ws.iter_rows(min_row=2):
            if row[review_col - 1].value is True:
                for cell in row:
                    cell.fill = REVIEW_FILL


def write_excel(results: list[Result], path: Path) -> Path:
    inv, lines = invoices_frame(results), lines_frame(results)
    summary = pd.DataFrame([
        {"metric": "documents processed", "value": len(results)},
        {"metric": "auto-accepted (no issues)", "value": int((~inv["needs_review"]).sum()) if len(inv) else 0},
        {"metric": "needs review", "value": int(inv["needs_review"].sum()) if len(inv) else 0},
        {"metric": "scanned (OCR)", "value": int(inv["scanned"].sum()) if len(inv) else 0},
        {"metric": "extracted by AI step", "value": int((inv["method"] == "ai").sum()) if len(inv) else 0},
    ])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        inv.to_excel(xw, sheet_name="Invoices", index=False)
        lines.to_excel(xw, sheet_name="Line items", index=False)
        summary.to_excel(xw, sheet_name="Summary", index=False)
        _style_sheet(xw.sheets["Invoices"], review_col=list(inv.columns).index("needs_review") + 1)
        _style_sheet(xw.sheets["Line items"])
        _style_sheet(xw.sheets["Summary"])
    return path


def write_csv(results: list[Result], out_dir: Path) -> list[Path]:
    a, b = out_dir / "invoices.csv", out_dir / "line_items.csv"
    invoices_frame(results).to_csv(a, index=False)
    lines_frame(results).to_csv(b, index=False)
    return [a, b]


def results_json(results: list[Result]) -> list[dict]:
    return [
        {"file": r.file, "method": r.method, "layout": r.layout, "scanned": r.scanned,
         "ocr_confidence": r.ocr_confidence, "needs_review": r.needs_review,
         "issues": [{"code": i.code, "message": i.message} for i in r.issues],
         "invoice": r.invoice.to_dict()}
        for r in results
    ]


def write_json(results: list[Result], path: Path) -> Path:
    path.write_text(json.dumps(results_json(results), indent=2, ensure_ascii=False))
    return path
